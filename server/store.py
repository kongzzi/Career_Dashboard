"""엑셀 파일 하나를 원본(single source of truth)으로 읽고 쓴다.

- 대시보드를 켜지 않아도 엑셀을 직접 열어 편집할 수 있게, 서식/드롭다운/틀고정을 넣어 저장한다.
- 저장할 때마다 직전 파일을 data/backups/ 로 옮겨 둔다.
"""
from __future__ import annotations

import datetime as dt
import os
import re
import shutil
import threading
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from schema import (
    ASSIGNMENT_FIELDS,
    CHOICES,
    FIELDS,
    FIELD_BY_KEY,
    LABEL_TO_KEY,
    ORG_FIELDS,
    SETTINGS,
    PAYROLL_FIELDS,
    SHEET_ASSIGN,
    SHEET_CONFIG,
    SHEET_ORGS,
    SHEET_PAYROLL,
    SHEET_RECORDS,
)

ROOT = Path(__file__).resolve().parent.parent

# 데이터 파일 위치. 기본은 data/career.xlsx 이고 CAREER_XLSX 로 바꿀 수 있다
# (run.py --data). 데모 데이터로 스크린샷을 찍거나, 실데이터를 동기화 폴더에
# 둘 때 쓴다. 백업은 언제나 이 파일 옆의 backups/ 로 간다.
XLSX_PATH = Path(os.environ.get("CAREER_XLSX") or ROOT / "data" / "career.xlsx").expanduser()
DATA_DIR = XLSX_PATH.parent
BACKUP_DIR = DATA_DIR / "backups"
MAX_BACKUPS = 10

_lock = threading.Lock()

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=12)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(bottom=THIN)


# --------------------------------------------------------------------------- 값 변환

def _to_date_str(v) -> str:
    """엑셀 셀 값을 'YYYY-MM-DD' 문자열로."""
    if v is None or v == "":
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    if not s:
        return ""
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return s


def _parse_date(s: str):
    s = (s or "").strip()
    if not s:
        return None
    try:
        return dt.date.fromisoformat(s)
    except ValueError:
        return None


def _to_time_str(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%H:%M")
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    s = str(v).strip()
    if not s:
        return ""
    for fmt in ("%H:%M:%S", "%H:%M", "%H시 %M분"):
        try:
            return dt.datetime.strptime(s, fmt).strftime("%H:%M")
        except ValueError:
            continue
    return s


def _to_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("원", "").strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


_FULL_DATE = re.compile(r"^(\d{4})[-./]\s*(\d{1,2})[-./]\s*(\d{1,2})")
_MD = re.compile(r"^(\d{1,2})[-./]\s*(\d{1,2})$")
_DAY = re.compile(r"^(\d{1,2})일?$")


def parse_date_list(raw, fallback_year: int | None = None) -> list[dt.date]:
    """'2026-08-10, 08-17, 24' 처럼 연·월을 이어받는 축약 표기를 허용해 날짜 목록을 읽는다.

    웹 화면의 parseDateList()와 같은 규칙 — 엑셀에서 직접 적어 넣어도 똑같이 해석된다.
    """
    if not raw:
        return []
    year = fallback_year or dt.date.today().year
    month = None
    out: list[dt.date] = []
    for tok in re.split(r"[,;\n]+|\s{2,}", str(raw)):
        t = tok.strip()
        if not t:
            continue
        m = _FULL_DATE.match(t)
        if m:
            year, month = int(m.group(1)), int(m.group(2))
            day = int(m.group(3))
        else:
            m = _MD.match(t)
            if m:
                month, day = int(m.group(1)), int(m.group(2))
            else:
                m = _DAY.match(t)
                if not m or month is None:
                    continue
                day = int(m.group(1))
        try:
            out.append(dt.date(year, month, day))
        except ValueError:
            continue
    return sorted(set(out))


def normalize_span(rec: dict) -> dict:
    """개별 날짜를 적었으면 그것이 곧 일정이다 — 시작일·종료일과 표기를 맞춰 둔다.

    엑셀만 열어 봐도 기간이 보이고, 날짜순 정렬이 정상 동작한다.
    """
    fallback = None
    start = _parse_date(str(rec.get("start_date") or ""))
    if start:
        fallback = start.year
    days = parse_date_list(rec.get("dates"), fallback)
    if not days:
        return rec
    rec["dates"] = ",".join(d.isoformat() for d in days)
    rec["start_date"] = days[0].isoformat()
    rec["end_date"] = days[-1].isoformat() if len(days) > 1 else ""
    return rec


def months_between(start: dt.date, end: dt.date | None) -> int:
    """계약 기간이 걸치는 달 수 (월정액 총액의 기준)"""
    if not start:
        return 0
    if not end or end <= start:
        return 1
    return max(1, (end.year - start.year) * 12 + (end.month - start.month) + 1)


def normalize_amounts(rec: dict) -> dict:
    """총액이 비어 있으면 급여방식에 따라 채운다 — 엑셀만 열어 봐도 금액이 보이도록.

    직접 적어 넣은 총액은 건드리지 않는다.
    """
    if _to_num(rec.get("gross")):
        return rec
    pay_type = (rec.get("pay_type") or "시간급").strip()
    if pay_type == "월정액(연봉제)":
        start = _parse_date(str(rec.get("start_date") or ""))
        end = _parse_date(str(rec.get("end_date") or ""))
        monthly = _to_num(rec.get("monthly_pay")) or 0
        if not monthly:
            annual = _to_num(rec.get("annual_pay")) or 0
            monthly = annual / 12 if annual else 0
        total = monthly * months_between(start, end) if start else 0
    elif pay_type == "총액 일괄":
        total = 0     # 총액을 직접 넣는 방식이라 계산하지 않는다
    else:
        total = (_to_num(rec.get("hours")) or 0) * (_to_num(rec.get("rate")) or 0)
    if total:
        rec["gross"] = round(total)
    return rec


def month_list(start: dt.date, end: dt.date | None) -> list[str]:
    """계약 기간이 걸치는 달들을 'YYYY-MM'으로"""
    if not start:
        return []
    last = end or start
    out, y, m = [], start.year, start.month
    while (y, m) <= (last.year, last.month) and len(out) < 60:
        out.append(f"{y:04d}-{m:02d}")
        m += 1
        if m > 12:
            m, y = 1, y + 1
    return out


def generate_payroll(rec: dict) -> list[dict]:
    """월정액 계약에서 달별 급여 행을 만든다. 금액은 예정액만 채우고 실수령은 비워 둔다."""
    start = _parse_date(str(rec.get("start_date") or ""))
    end = _parse_date(str(rec.get("end_date") or ""))
    months = month_list(start, end)
    if not months:
        return []
    monthly = _to_num(rec.get("monthly_pay")) or 0
    if not monthly:
        annual = _to_num(rec.get("annual_pay")) or 0
        monthly = annual / 12 if annual else 0
    if not monthly:
        gross = _to_num(rec.get("gross")) or 0
        monthly = gross / len(months) if gross else 0
    return [{
        "record_id": rec.get("id", ""),
        "month": m,
        "gross": round(monthly) if monthly else None,
        "deduction": None, "net": None, "paid_date": "",
        "status": "예정", "memo": "",
    } for m in months]


def _cell_out(field, raw):
    """엑셀 셀 -> JSON 값"""
    t = field["type"]
    if t == "date":
        return _to_date_str(raw)
    if t == "time":
        return _to_time_str(raw)
    if t in ("number", "money"):
        n = _to_num(raw)
        return n
    return "" if raw is None else str(raw).strip()


def _cell_in(field, value):
    """JSON 값 -> 엑셀 셀"""
    t = field["type"]
    if t == "date":
        return _parse_date(value if isinstance(value, str) else "")
    if t in ("number", "money"):
        n = _to_num(value)
        return n
    if value is None:
        return None
    s = str(value).strip()
    return s or None


# --------------------------------------------------------------------------- 시트 생성

def _style_header(ws, labels, widths):
    for i, (label, width) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=1, column=i, value=label)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"


def _apply_validations(ws, n_rows=2000, fmt_rows=60):
    """드롭다운 + 숫자/날짜 표시형식.

    드롭다운은 범위만 지정하면 되지만 표시형식은 셀을 실제로 만들어야 하므로,
    데이터가 있는 구간 + 여유분에만 적용한다(빈 셀 수천 개로 파일이 붓는 것을 막는다).
    """
    for i, f in enumerate(FIELDS, start=1):
        col = get_column_letter(i)
        rng = f"{col}2:{col}{n_rows}"
        if f["type"] == "select":
            opts = CHOICES.get(f.get("options"), [])
            if opts:
                dv = DataValidation(
                    type="list",
                    formula1='"' + ",".join(opts) + '"',
                    allow_blank=True,
                    showDropDown=False,
                )
                ws.add_data_validation(dv)
                dv.add(rng)
        fmt = None
        if f["type"] == "money":
            fmt = "#,##0"
        elif f["type"] == "date":
            fmt = "yyyy-mm-dd"
        elif f["type"] == "number":
            fmt = "0.##"
        if fmt:
            for r in range(2, fmt_rows + 2):
                ws.cell(row=r, column=i).number_format = fmt


def create_workbook(path: Path = XLSX_PATH) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = SHEET_RECORDS
    _style_header(ws, [f["label"] for f in FIELDS], [f["width"] for f in FIELDS])
    _apply_validations(ws)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(FIELDS))}1"

    wo = wb.create_sheet(SHEET_ORGS)
    _style_header(wo, [f["label"] for f in ORG_FIELDS], [f["width"] for f in ORG_FIELDS])

    wa = wb.create_sheet(SHEET_ASSIGN)
    _style_header(wa, [f["label"] for f in ASSIGNMENT_FIELDS], [f["width"] for f in ASSIGNMENT_FIELDS])

    wp = wb.create_sheet(SHEET_PAYROLL)
    _style_header(wp, [f["label"] for f in PAYROLL_FIELDS], [f["width"] for f in PAYROLL_FIELDS])

    wc = wb.create_sheet(SHEET_CONFIG)
    _style_header(wc, ["구분", "키", "값", "설명"], [16, 20, 22, 40])
    r = 2
    for s in SETTINGS:
        wc.cell(row=r, column=1, value="설정")
        wc.cell(row=r, column=2, value=s["label"])
        wc.cell(row=r, column=3, value=s["value"])
        wc.cell(row=r, column=4, value=s["note"])
        r += 1
    r += 1
    wc.cell(row=r, column=1, value="※ 아래 목록은 대시보드 드롭다운입니다. 행을 추가하면 선택지가 늘어납니다.").font = TITLE_FONT
    r += 1
    for name, opts in CHOICES.items():
        for o in opts:
            wc.cell(row=r, column=1, value="목록")
            wc.cell(row=r, column=2, value=name)
            wc.cell(row=r, column=3, value=o)
            r += 1

    wb.save(path)
    return path


def ensure_workbook(path: Path = XLSX_PATH) -> Path:
    if not path.exists():
        create_workbook(path)
    return path


# --------------------------------------------------------------------------- 읽기

def _row_is_empty(values) -> bool:
    return all(v is None or str(v).strip() == "" for v in values)


def _read_sheet(ws, fields) -> list[dict]:
    """헤더 이름으로 컬럼을 찾아 읽는다.

    위치로 읽으면 나중에 컬럼을 중간에 하나 끼워 넣는 순간 예전 파일의 값이 통째로 밀린다.
    (일감 시트가 처음부터 이 방식이라 스키마가 여러 번 바뀌어도 멀쩡했다.)
    """
    headers = [c.value for c in ws[1]]
    col_of = {}
    for idx, h in enumerate(headers):
        if h is None:
            continue
        label = str(h).strip()
        for f in fields:
            if f["label"] == label:
                col_of[f["key"]] = idx
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if _row_is_empty(row):
            continue
        item = {}
        for f in fields:
            idx = col_of.get(f["key"])
            item[f["key"]] = _cell_out(f, row[idx] if idx is not None and idx < len(row) else None)
        out.append(item)
    return out


def read_all(path: Path = XLSX_PATH) -> dict:
    ensure_workbook(path)
    wb = load_workbook(path, data_only=True)

    records = []
    if SHEET_RECORDS in wb.sheetnames:
        ws = wb[SHEET_RECORDS]
        headers = [c.value for c in ws[1]]
        col_of = {}
        for idx, h in enumerate(headers):
            if h is None:
                continue
            key = LABEL_TO_KEY.get(str(h).strip())
            if key:
                col_of[key] = idx
        for row in ws.iter_rows(min_row=2, values_only=True):
            if _row_is_empty(row):
                continue
            rec = {}
            for f in FIELDS:
                idx = col_of.get(f["key"])
                raw = row[idx] if idx is not None and idx < len(row) else None
                rec[f["key"]] = _cell_out(f, raw)
            records.append(rec)

    orgs = _read_sheet(wb[SHEET_ORGS], ORG_FIELDS) if SHEET_ORGS in wb.sheetnames else []

    assignments = [a for a in (_read_sheet(wb[SHEET_ASSIGN], ASSIGNMENT_FIELDS)
                               if SHEET_ASSIGN in wb.sheetnames else [])
                   if a.get("record_id")]

    payroll = [p for p in (_read_sheet(wb[SHEET_PAYROLL], PAYROLL_FIELDS)
                           if SHEET_PAYROLL in wb.sheetnames else [])
               if p.get("record_id")]

    settings = {s["key"]: s["value"] for s in SETTINGS}
    label_to_setting = {s["label"]: s["key"] for s in SETTINGS}
    choices = {k: list(v) for k, v in CHOICES.items()}
    if SHEET_CONFIG in wb.sheetnames:
        found_lists: dict[str, list] = {}
        for row in wb[SHEET_CONFIG].iter_rows(min_row=2, values_only=True):
            kind = (row[0] or "") if len(row) > 0 else ""
            key = (row[1] or "") if len(row) > 1 else ""
            val = row[2] if len(row) > 2 else None
            kind, key = str(kind).strip(), str(key).strip()
            if kind == "설정" and key in label_to_setting:
                n = _to_num(val)
                settings[label_to_setting[key]] = n if n is not None else 0
            elif kind == "목록" and key and val not in (None, ""):
                found_lists.setdefault(key, []).append(str(val).strip())
        # 시트에 적힌 목록이 우선이되, 나중에 코드에 추가된 선택지는 뒤에 붙인다.
        # (예전 파일이 새 기능의 선택지를 가리는 것을 막는다)
        for k, v in found_lists.items():
            if not v:
                continue
            merged = list(v)
            for d in CHOICES.get(k, []):
                if d not in merged:
                    merged.append(d)
            choices[k] = merged

    # 신규 ID 채번용
    max_id = 0
    for r in records:
        s = str(r.get("id") or "")
        if s.upper().startswith("R"):
            try:
                max_id = max(max_id, int(s[1:]))
            except ValueError:
                pass
    return {
        "records": records,
        "assignments": assignments,
        "payroll": payroll,
        "orgs": orgs,
        "settings": settings,
        "choices": choices,
        "next_id": f"R{max_id + 1:04d}",
        "file": str(path),
        "mtime": path.stat().st_mtime,
    }


# --------------------------------------------------------------------------- 쓰기

def _backup(path: Path):
    if not path.exists():
        return
    # 백업은 언제나 원본 옆에 만든다. path 를 바꿔 부르면(테스트 등) 백업도 따라가야
    # 실데이터의 백업 이력이 엉뚱한 파일로 밀려나지 않는다.
    # 기본 경로에서는 결과가 BACKUP_DIR 과 같다.
    backup_dir = path.parent / BACKUP_DIR.name
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    shutil.copy2(path, backup_dir / f"{path.stem}-{stamp}{path.suffix}")
    files = sorted(backup_dir.glob(f"{path.stem}-*{path.suffix}"))
    for old in files[:-MAX_BACKUPS]:
        old.unlink(missing_ok=True)


def write_records(records: list[dict], orgs: list[dict] | None = None,
                  settings: dict | None = None, payroll: list[dict] | None = None,
                  assignments: list[dict] | None = None, path: Path = XLSX_PATH) -> dict:
    """레코드 전체를 다시 쓴다. 행 수가 수백 단위라 전체 재작성이 가장 안전하다."""
    with _lock:
        ensure_workbook(path)
        _backup(path)
        wb = load_workbook(path)

        if SHEET_RECORDS in wb.sheetnames:
            del wb[SHEET_RECORDS]
        ws = wb.create_sheet(SHEET_RECORDS, 0)
        _style_header(ws, [f["label"] for f in FIELDS], [f["width"] for f in FIELDS])
        n_rows = max(2000, len(records) + 200)
        _apply_validations(ws, n_rows, fmt_rows=len(records) + 40)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(FIELDS))}1"

        for r_i, rec in enumerate(records, start=2):
            for c_i, f in enumerate(FIELDS, start=1):
                ws.cell(row=r_i, column=c_i, value=_cell_in(f, rec.get(f["key"])))
            ws.row_dimensions[r_i].height = 18

        if assignments is not None:
            if SHEET_ASSIGN in wb.sheetnames:
                del wb[SHEET_ASSIGN]
            wa = wb.create_sheet(SHEET_ASSIGN)
            _style_header(wa, [f["label"] for f in ASSIGNMENT_FIELDS], [f["width"] for f in ASSIGNMENT_FIELDS])
            rows = sorted(assignments, key=lambda x: (str(x.get("record_id") or ""),
                                                      str(x.get("subject") or ""), str(x.get("section") or "")))
            for r_i, item in enumerate(rows, start=2):
                for c_i, f in enumerate(ASSIGNMENT_FIELDS, start=1):
                    wa.cell(row=r_i, column=c_i, value=_cell_in(f, item.get(f["key"])))

        if payroll is not None:
            if SHEET_PAYROLL in wb.sheetnames:
                del wb[SHEET_PAYROLL]
            wp = wb.create_sheet(SHEET_PAYROLL)
            _style_header(wp, [f["label"] for f in PAYROLL_FIELDS], [f["width"] for f in PAYROLL_FIELDS])
            rows = sorted(payroll, key=lambda x: (str(x.get("record_id") or ""), str(x.get("month") or "")))
            for r_i, item in enumerate(rows, start=2):
                for c_i, f in enumerate(PAYROLL_FIELDS, start=1):
                    wp.cell(row=r_i, column=c_i, value=_cell_in(f, item.get(f["key"])))
            for c_i, f in enumerate(PAYROLL_FIELDS, start=1):
                fmt = {"money": "#,##0", "date": "yyyy-mm-dd"}.get(f["type"])
                if not fmt:
                    continue
                for r_i in range(2, len(rows) + 2):
                    wp.cell(row=r_i, column=c_i).number_format = fmt

        if orgs is not None:
            if SHEET_ORGS in wb.sheetnames:
                del wb[SHEET_ORGS]
            wo = wb.create_sheet(SHEET_ORGS)
            _style_header(wo, [f["label"] for f in ORG_FIELDS], [f["width"] for f in ORG_FIELDS])
            for r_i, o in enumerate(orgs, start=2):
                for c_i, f in enumerate(ORG_FIELDS, start=1):
                    wo.cell(row=r_i, column=c_i, value=_cell_in(f, o.get(f["key"])))

        if settings:
            wc = wb[SHEET_CONFIG] if SHEET_CONFIG in wb.sheetnames else wb.create_sheet(SHEET_CONFIG)
            label_of = {s["key"]: s["label"] for s in SETTINGS}
            wanted = {label_of[k]: v for k, v in settings.items() if k in label_of}
            seen, last_row = set(), 1
            for row in wc.iter_rows(min_row=2):
                if str(row[0].value or "").strip() == "설정":
                    label = str(row[1].value or "").strip()
                    last_row = row[0].row
                    seen.add(label)
                    if label in wanted:
                        row[2].value = wanted[label]

            # 코드에 새로 생긴 드롭다운 선택지도 시트 끝에 붙여 파일이 따라오게 한다
            existing_opts: dict[str, set] = {}
            for row in wc.iter_rows(min_row=2, values_only=True):
                if str(row[0] or "").strip() == "목록" and len(row) > 2:
                    existing_opts.setdefault(str(row[1] or "").strip(), set()).add(str(row[2] or "").strip())
            add_opts = [
                (name, o) for name, opts in CHOICES.items() for o in opts
                if name in existing_opts and o not in existing_opts[name]
            ]
            for name, o in add_opts:
                r = wc.max_row + 1
                wc.cell(row=r, column=1, value="목록")
                wc.cell(row=r, column=2, value=name)
                wc.cell(row=r, column=3, value=o)

            # 나중에 추가된 설정 항목은 예전 파일에 행 자체가 없다 — 만들어 넣는다
            note_of = {s["label"]: s["note"] for s in SETTINGS}
            missing = [(lab, v) for lab, v in wanted.items() if lab not in seen]
            if missing:
                wc.insert_rows(last_row + 1, amount=len(missing))
                for i, (lab, v) in enumerate(missing):
                    r = last_row + 1 + i
                    wc.cell(row=r, column=1, value="설정")
                    wc.cell(row=r, column=2, value=lab)
                    wc.cell(row=r, column=3, value=v)
                    wc.cell(row=r, column=4, value=note_of.get(lab, ""))

        wb.save(path)
    return read_all(path)
